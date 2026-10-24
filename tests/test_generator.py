"""Tests for the data generator.

Run: uv run pytest   (or: python -m unittest discover -s tests)
They use a short date range, so they finish in a few seconds.
"""

import csv
import json
import tempfile
import unittest
from datetime import date
from pathlib import Path

from data_generator.generate import Generator, State, fiscal_year, main
from data_generator.writers import API_PAGE_SIZE


def read_csv(path: Path):
    with path.open(encoding="utf-8") as f:
        return list(csv.DictReader(f))


class FiscalYearTests(unittest.TestCase):
    def test_fiscal_year_boundaries(self):
        self.assertEqual(fiscal_year(date(2025, 6, 30)), 2025)
        self.assertEqual(fiscal_year(date(2025, 7, 1)), 2026)


class GeneratorTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = Path(tempfile.mkdtemp())
        cls.out, cls.state = cls.tmp / "data", cls.tmp / "state.json"
        main(["full", "--start", "2025-07-01", "--end", "2025-12-31", "--out", str(cls.out), "--state", str(cls.state)])
        main(["incremental", "--days", "2", "--out", str(cls.out), "--state", str(cls.state)])

    def _rows(self, source, entity, load_date):
        folder = {"tessitura": "tessitura_db", "marketing": "marketing_adls", "mip": "mip_db"}[source]
        return read_csv(self.out / folder / entity / f"load_date={load_date}" / f"{entity}.csv")

    def _api(self, entity, load_date, api="tessitura_api"):
        """Follow next_link from page 1, the way the notebook does."""
        rows, page, count = [], "page-1.json", None
        while page:
            body = json.loads((self.out / api / "v1" / entity / load_date / page).read_text())
            count = body["count"]
            rows += body["value"]
            page = body["next_link"]
        return count, rows

    def test_partitions_and_manifests_exist(self):
        for d in ("2025-12-31", "2026-01-01", "2026-01-02"):
            for folder, n in (("tessitura_db", 7), ("marketing_adls", 3), ("mip_db", 5)):
                m = json.loads((self.out / folder / "_manifests" / f"load_date={d}.json").read_text())
                self.assertEqual(len(m["files"]), n)

    def test_manifest_counts_match_files(self):
        m = json.loads((self.out / "tessitura_db" / "_manifests" / "load_date=2025-12-31.json").read_text())
        orders = next(f for f in m["files"] if f["entity"] == "order")
        self.assertEqual(orders["rows"], len(self._rows("tessitura", "order", "2025-12-31")))

    def test_api_paging_and_index(self):
        index = json.loads((self.out / "tessitura_api" / "v1" / "index.json").read_text())
        self.assertEqual(index["load_dates"], ["2025-12-31", "2026-01-01", "2026-01-02"])
        count, rows = self._api("contribution", "2025-12-31")
        self.assertGreater(count, API_PAGE_SIZE, "expected more than one page")
        self.assertEqual(len({r["contribution_id"] for r in rows}), count)
        self.assertGreater(len(rows), count, "expected overlapping pages on the first load")

    def test_education_workbooks(self):
        from openpyxl import load_workbook
        ref = load_workbook(self.out / "education_sharepoint" / "Program Reference.xlsx")
        self.assertEqual(ref.sheetnames, ["Programs", "Schools"])
        sessions = sorted((self.out / "education_sharepoint" / "Sessions").glob("*.xlsx"))
        self.assertTrue(sessions)
        header = next(load_workbook(sessions[0])["Sessions"].iter_rows(max_row=1, values_only=True))
        self.assertEqual(header[0], "Enrollment ID")

    def test_incremental_ids_do_not_collide_with_history(self):
        hist = {r["order_id"] for r in self._rows("tessitura", "order", "2025-12-31")}
        inc = [r for r in self._rows("tessitura", "order", "2026-01-01") if r["status"] == "Complete"]
        self.assertTrue(inc, "incremental produced no new orders")
        self.assertFalse(hist & {r["order_id"] for r in inc})

    def test_defects_are_injected(self):
        clicks = self._rows("marketing", "email_clicks", "2025-12-31")
        ids = [r["click_id"] for r in clicks]
        self.assertGreater(len(ids), len(set(ids)), "expected clicks exported twice")
        custs = self._rows("tessitura", "constituent", "2025-12-31")
        self.assertTrue(any(c["email"] == "" for c in custs), "expected missing emails")

    def test_never_oversold(self):
        perfs = {p["performance_id"]: int(p["capacity"]) for p in self._rows("tessitura", "performance", "2025-12-31")}
        sold = {}
        for r in self._rows("tessitura", "order_line", "2025-12-31"):
            if r["performance_id"] not in perfs:
                continue
            sold[r["performance_id"]] = sold.get(r["performance_id"], 0) + int(r["quantity"])
        self.assertTrue(all(sold[p] <= perfs[p] for p in sold))

    def test_contributions_use_constituent_ids(self):
        """Donors are constituents: every contribution points at a constituent record
        (apart from the planted unknown ID), and some donors have a second, duplicate record."""
        constituents = {r["constituent_id"] for d in ("2025-12-31", "2026-01-01", "2026-01-02")
                        for r in self._rows("tessitura", "constituent", d)}
        _, rows = self._api("contribution", "2025-12-31")
        unknown = {r["constituent_id"] for r in rows} - constituents
        self.assertLessEqual(unknown, {"T9999999"})
        self.assertTrue(any(c.startswith("T5") for c in constituents), "expected donor-created constituent records")

    def test_ad_spend_restates_last_two_days(self):
        days = {r["ad_date"] for r in self._rows("marketing", "ad_spend_daily", "2026-01-02")}
        self.assertEqual(days, {"2025-12-31", "2026-01-01", "2026-01-02"})

    def test_cash_payments_and_budget(self):
        """Payroll and bills are paid from cash after month end, and every fiscal year has a budget."""
        rows = self._rows("mip", "gl_transaction", "2025-12-31")
        modules = {r["source_module"] for r in rows}
        self.assertTrue({"OPENING", "PAYROLL_PAYMENT", "AP_PAYMENT"} <= modules)
        cash = sum(float(r["debit_amount"]) - float(r["credit_amount"]) for r in rows if r["account_code"] == "1010")
        self.assertGreater(cash, 0)
        budget = self._rows("mip", "gl_budget", "2025-12-31")
        self.assertEqual({r["fiscal_year"] for r in budget}, {"2026"})
        self.assertTrue(all(float(r["amount"]) > 0 for r in budget))

    def test_every_ledger_journal_balances(self):
        for d in ("2025-12-31", "2026-01-01", "2026-01-02"):
            totals = {}
            for r in self._rows("mip", "gl_transaction", d):
                totals[r["journal_id"]] = totals.get(r["journal_id"], 0) + float(r["credit_amount"]) - float(r["debit_amount"])
            self.assertTrue(totals, f"no journals on {d}")
            self.assertTrue(all(abs(v) < 0.01 for v in totals.values()), f"unbalanced journal on {d}")

    def test_rental_revenue_ties_to_diese_except_planted(self):
        """The ledger's rental revenue equals Diese's rentals, plus the few planted ones that
        were cancelled after invoicing (so the ledger can only be higher)."""
        _, events = self._api("event", "2025-12-31", "diese_api")
        _, bookings = self._api("booking", "2025-12-31", "diese_api")
        ev = {e["event_id"]: e for e in events}
        bookings = list({b["booking_id"]: b for b in bookings}.values())   # pages overlap; Silver dedupes the same way
        diese = sum(b["rental_fee"] for b in bookings if ev[b["event_id"]]["event_type"] == "Rental"
                    and ev[b["event_id"]]["status"] != "Cancelled" and ev[b["event_id"]]["event_date"] <= "2025-12-31")
        gl = sum(float(r["credit_amount"]) - float(r["debit_amount"])
                 for r in self._rows("mip", "gl_transaction", "2025-12-31") if r["account_code"] == "4300")
        self.assertGreater(diese, 0)
        self.assertGreaterEqual(round(gl, 2), round(diese, 2))
        self.assertLess(gl - diese, 0.05 * diese)

    def test_concert_events_link_to_tessitura(self):
        _, events = self._api("event", "2025-12-31", "diese_api")
        perf_ids = {p["performance_id"] for p in self._rows("tessitura", "performance", "2025-12-31")}
        concerts = [e for e in events if e["event_type"] == "Concert"]
        linked = [e for e in concerts if e["tessitura_performance_id"]]
        self.assertTrue(all(e["tessitura_performance_id"] in perf_ids for e in linked))
        self.assertGreater(len(linked), 0.95 * len(concerts))

    def test_deterministic_for_same_seed(self):
        a = Generator(State(seed=7, start="2025-09-01", end="2025-09-30")).full()[1]
        b = Generator(State(seed=7, start="2025-09-01", end="2025-09-30")).full()[1]
        self.assertEqual(len(a[("tessitura_db", "order")]), len(b[("tessitura_db", "order")]))
        self.assertEqual(a[("tessitura_api", "contribution")][:5], b[("tessitura_api", "contribution")][:5])


if __name__ == "__main__":
    unittest.main()
