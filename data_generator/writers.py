"""Writes each source system's data in the shape that system really hands it over.

  ticketing    -> data/ticketing_db/<table>/load_date=D/<table>.csv     (loaded into SQL Server)
  fundraising  -> data/fundraising_api/v1/<table>/D/page-N.json       (a static JSON API)
  education    -> data/education_sharepoint/*.xlsx                     (spreadsheets for SharePoint)
  marketing    -> data/marketing_adls/<table>/load_date=D/<table>.csv  (exports for ADLS Gen2)
"""

from __future__ import annotations

import csv
import json
from datetime import date, datetime
from pathlib import Path

FOLDERS = {
    "ticketing": "ticketing_db",
    "fundraising": "fundraising_api",
    "education": "education_sharepoint",
    "marketing": "marketing_adls",
}
API_PAGE_SIZE = 500


def write_all(load_date: date, batch: dict, out: Path, full: bool) -> dict:
    by_source: dict[str, dict] = {}
    for (source, entity), rows in batch.items():
        by_source.setdefault(source, {})[entity] = rows
    files = []
    files += _csv_partitions(out / FOLDERS["ticketing"], "ticketing", load_date, by_source["ticketing"])
    files += _api_pages(out / FOLDERS["fundraising"], load_date, by_source["fundraising"], full)
    files += _workbooks(out / FOLDERS["education"], by_source["education"])
    files += _csv_partitions(out / FOLDERS["marketing"], "marketing", load_date, by_source["marketing"])
    return {"load_date": load_date.isoformat(), "files": files}


# ----------------------------------------------------------------------------- CSV folders
def _csv_partitions(root: Path, source: str, load_date: date, tables: dict) -> list[dict]:
    """One CSV per table per day, plus a manifest with the row counts."""
    d = load_date.isoformat()
    files = []
    for entity, rows in tables.items():
        folder = root / entity / f"load_date={d}"
        folder.mkdir(parents=True, exist_ok=True)
        path = folder / f"{entity}.csv"
        with path.open("w", newline="", encoding="utf-8") as f:
            if rows:
                w = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
                w.writeheader()
                w.writerows(rows)
        files.append({"source": source, "entity": entity, "rows": len(rows),
                      "path": path.relative_to(root).as_posix(), "where": f"{root.name}/{entity}/"})
    manifest = {"load_date": d, "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                "files": [{k: f[k] for k in ("source", "entity", "path", "rows")} for f in files]}
    mpath = root / "_manifests" / f"load_date={d}.json"
    mpath.parent.mkdir(parents=True, exist_ok=True)
    mpath.write_text(json.dumps(manifest, indent=2))
    return files


# ----------------------------------------------------------------------------- JSON API
def _api_pages(root: Path, load_date: date, tables: dict, full: bool) -> list[dict]:
    """Pages of 500 records, each pointing to the next one, like most SaaS APIs (Blackbaud SKY API uses
    the same count / value / next_link shape). Bad data: on the big first load, each page repeats the
    last record of the page before, which is what offset paging does when records are added mid-read."""
    d = load_date.isoformat()
    files = []
    for entity, rows in tables.items():
        folder = root / "v1" / entity / d
        folder.mkdir(parents=True, exist_ok=True)
        chunks = [rows[i:i + API_PAGE_SIZE] for i in range(0, len(rows), API_PAGE_SIZE)] or [[]]
        for n, chunk in enumerate(chunks, start=1):
            if full and n > 1:
                chunk = [chunks[n - 2][-1]] + chunk
            page = {"count": len(rows), "value": chunk, "next_link": f"page-{n + 1}.json" if n < len(chunks) else None}
            (folder / f"page-{n}.json").write_text(json.dumps(page, separators=(",", ":")))
        files.append({"source": "fundraising", "entity": entity, "rows": len(rows),
                      "where": f"{root.name}/v1/{entity}/{d}/ ({len(chunks)} page(s))"})

    # index.json tells the notebook which days exist; a static site can't filter by date itself.
    v1 = root / "v1"
    load_dates = sorted({p.name for e in v1.iterdir() if e.is_dir() for p in e.iterdir() if p.is_dir()})
    index = {"api": "Harmonia Hall fundraising API (synthetic data)", "version": "v1", "page_size": API_PAGE_SIZE,
             "entities": sorted(e.name for e in v1.iterdir() if e.is_dir()), "load_dates": load_dates,
             "first_page": "{entity}/{load_date}/page-1.json"}
    (v1 / "index.json").write_text(json.dumps(index, indent=2))
    (root / ".nojekyll").write_text("")              # tells GitHub Pages to serve the files as they are
    (root / "README.md").write_text(
        "# Harmonia Hall fundraising API (mock)\n\nSynthetic data for a Microsoft Fabric practice project. "
        "Start at `v1/index.json`.\nNo real people or donations.\n")
    return files


# ----------------------------------------------------------------------------- Excel
PROGRAM_HEADERS = {"program_id": "Program ID", "program_name": "Program Name", "program_type": "Program Type", "audience": "Audience"}
SCHOOL_HEADERS = {"school_id": "School ID", "school_name": "School Name", "borough": "Borough", "is_title_i": "Title I"}
SESSION_HEADERS = {"enrollment_id": "Enrollment ID", "program_id": "Program ID", "school_id": "School ID",
                   "session_date": "Session Date", "participants": "Participants", "teaching_artist_hours": "Teaching Artist Hours"}


def _workbooks(root: Path, tables: dict) -> list[dict]:
    """The way a program team keeps its records: a reference workbook, plus one sessions workbook per month.
    Headers are written for people (Program ID, Title I = Yes/No); the Dataflow renames and types them."""
    from openpyxl import Workbook, load_workbook

    root.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Programs"
    _sheet(ws, PROGRAM_HEADERS, tables["programs"])
    _sheet(wb.create_sheet("Schools"), SCHOOL_HEADERS,
           [{**s, "is_title_i": "Yes" if s["is_title_i"] else "No"} for s in tables["schools"]])
    wb.save(root / "Program Reference.xlsx")

    sessions_dir = root / "Sessions"
    sessions_dir.mkdir(exist_ok=True)
    by_month: dict[str, list] = {}
    for r in tables["enrollments"]:
        by_month.setdefault(r["session_date"][:7], []).append(r)
    for month, rows in by_month.items():
        path = sessions_dir / f"Sessions {month}.xlsx"
        if path.exists():                                     # add the new days to this month's workbook
            wb = load_workbook(path)
            ws = wb["Sessions"]
            have = {row[0] for row in ws.iter_rows(min_row=2, max_col=1, values_only=True)}
            for r in rows:
                if r["enrollment_id"] not in have:
                    ws.append(_session_cells(r))
        else:
            wb = Workbook()
            ws = wb.active
            ws.title = "Sessions"
            ws.append(list(SESSION_HEADERS.values()))
            for r in rows:
                ws.append(_session_cells(r))
            ws.column_dimensions["D"].width = 14
        for cell in ws["D"][1:]:
            cell.number_format = "yyyy-mm-dd"
        wb.save(path)

    return [{"source": "education", "entity": "programs", "rows": len(tables["programs"]),
             "where": f"{root.name}/Program Reference.xlsx"},
            {"source": "education", "entity": "schools", "rows": len(tables["schools"]),
             "where": f"{root.name}/Program Reference.xlsx"},
            {"source": "education", "entity": "enrollments", "rows": len(tables["enrollments"]),
             "where": f"{root.name}/Sessions/ ({len(by_month)} workbook(s))"}]


def _sheet(ws, headers: dict, rows: list[dict]):
    ws.append(list(headers.values()))
    for r in rows:
        ws.append([r[k] for k in headers])


def _session_cells(r: dict) -> list:
    return [r["enrollment_id"], r["program_id"], r["school_id"] or None, date.fromisoformat(r["session_date"]),
            r["participants"], r["teaching_artist_hours"]]
